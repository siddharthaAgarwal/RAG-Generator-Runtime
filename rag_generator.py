#!/usr/bin/env python3
"""A small local RAG pipeline: ingest documents, build an index, and retrieve context.

Examples:
  python rag_generator.py index ./docs/report.pdf notes.txt --index-dir ./rag_index
  python rag_generator.py query "What does the report say about renewals?" --index-dir ./rag_index
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Iterable

import faiss
import numpy as np
from docx import Document
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer


MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"
SUPPORTED_SUFFIXES = {".pdf", ".txt", ".docx"}
INDEX_FILE = "index.faiss"
CHUNKS_FILE = "chunks.json"
CONFIG_FILE = "config.json"


def collect_files(inputs: list[str]) -> list[Path]:
    """Return supported files from paths and recursively from directories."""
    files: list[Path] = []
    for raw_path in inputs:
        path = Path(raw_path).expanduser()
        if not path.exists():
            print(f"Warning: path not found, skipping: {path}", file=sys.stderr)
            continue
        candidates: Iterable[Path] = path.rglob("*") if path.is_dir() else [path]
        files.extend(candidate for candidate in candidates if candidate.is_file() and candidate.suffix.lower() in SUPPORTED_SUFFIXES)
    return sorted(set(files))


def extract_text(path: Path) -> str:
    """Extract readable text from one supported document."""
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        return "\n".join(paragraphs + table_text)
    raise ValueError(f"Unsupported file type: {path}")


def split_sentences(text: str) -> list[str]:
    """Lightweight sentence splitting with no additional NLP download."""
    cleaned = re.sub(r"\s+", " ", text).strip()
    return [part.strip() for part in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9])", cleaned) if part.strip()]


def chunk_text(text: str, chunk_size: int = 900, overlap: int = 180) -> list[str]:
    """Create sentence-aware chunks, carrying a tail forward for context overlap."""
    if chunk_size <= overlap:
        raise ValueError("chunk_size must be greater than overlap")

    sentences = split_sentences(text)
    if not sentences:
        return []

    chunks: list[str] = []
    current = ""
    for sentence in sentences:
        # Split unusually long single sentences, preserving all content.
        pieces = [sentence[i : i + chunk_size] for i in range(0, len(sentence), chunk_size)]
        for piece in pieces:
            candidate = f"{current} {piece}".strip()
            if current and len(candidate) > chunk_size:
                chunks.append(current)
                current = current[-overlap:].lstrip() if overlap else ""
                current = f"{current} {piece}".strip()
            else:
                current = candidate
    if current:
        chunks.append(current)
    return chunks


def load_model() -> SentenceTransformer:
    print(f"Loading embedding model: {MODEL_NAME}")
    return SentenceTransformer(MODEL_NAME)


def build_index(file_paths: list[Path], index_dir: Path, chunk_size: int, overlap: int) -> None:
    records: list[dict[str, object]] = []
    for path in file_paths:
        try:
            text = extract_text(path)
            chunks = chunk_text(text, chunk_size, overlap)
            records.extend({"source": str(path.resolve()), "chunk_id": number, "text": chunk} for number, chunk in enumerate(chunks))
            print(f"{path}: {len(chunks)} chunks")
        except Exception as error:
            print(f"Warning: could not read {path}: {error}", file=sys.stderr)

    if not records:
        raise RuntimeError("No text chunks were extracted. Check that the documents contain selectable text.")

    model = load_model()
    embeddings = model.encode([str(record["text"]) for record in records], normalize_embeddings=True, show_progress_bar=True)
    vectors = np.ascontiguousarray(np.asarray(embeddings, dtype=np.float32))

    # Inner product on unit-normalized embeddings is cosine similarity.
    index = faiss.IndexHNSWFlat(vectors.shape[1], 32, faiss.METRIC_INNER_PRODUCT)
    index.hnsw.efConstruction = 80
    index.hnsw.efSearch = 64
    index.add(vectors)

    index_dir.mkdir(parents=True, exist_ok=True)
    faiss.write_index(index, str(index_dir / INDEX_FILE))
    (index_dir / CHUNKS_FILE).write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    (index_dir / CONFIG_FILE).write_text(json.dumps({"model": MODEL_NAME, "chunk_size": chunk_size, "overlap": overlap, "metric": "cosine similarity (normalized inner product)"}, indent=2), encoding="utf-8")
    print(f"Indexed {len(records)} chunks from {len(file_paths)} file(s) in {index_dir}")


def query_index(question: str, index_dir: Path, top_k: int) -> None:
    index_path, chunks_path = index_dir / INDEX_FILE, index_dir / CHUNKS_FILE
    if not index_path.exists() or not chunks_path.exists():
        raise FileNotFoundError(f"No saved index in {index_dir}. Run the 'index' command first.")

    records = json.loads(chunks_path.read_text(encoding="utf-8"))
    index = faiss.read_index(str(index_path))
    model = load_model()
    query_vector = model.encode([question], normalize_embeddings=True)
    scores, ids = index.search(np.ascontiguousarray(np.asarray(query_vector, dtype=np.float32)), min(top_k, len(records)))

    results = []
    for score, record_id in zip(scores[0], ids[0]):
        if record_id >= 0:
            record = records[int(record_id)]
            results.append((float(score), record))

    print("\nRetrieved chunks:")
    for rank, (score, record) in enumerate(results, start=1):
        print(f"\n[{rank}] cosine similarity: {score:.4f}")
        print(f"source: {record['source']} (chunk {record['chunk_id']})")
        print(record["text"])

    context = "\n\n---\n\n".join(
        f"Source: {record['source']} (chunk {record['chunk_id']})\n{record['text']}" for _, record in results
    )
    print("\nAnswer context (pass this to your preferred LLM):\n")
    print(context)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Local, plain-Python RAG retrieval with FAISS HNSW.")
    subparsers = parser.add_subparsers(dest="command", required=True)
    index_parser = subparsers.add_parser("index", help="Extract, chunk, embed, and persist documents")
    index_parser.add_argument("inputs", nargs="+", help="PDF, TXT, DOCX files or directories")
    index_parser.add_argument("--index-dir", type=Path, default=Path("rag_index"))
    index_parser.add_argument("--chunk-size", type=int, default=900, help="Approximate characters per semantic chunk")
    index_parser.add_argument("--overlap", type=int, default=180, help="Characters carried into the next chunk")
    query_parser = subparsers.add_parser("query", help="Retrieve relevant chunks and assemble answer context")
    query_parser.add_argument("question")
    query_parser.add_argument("--index-dir", type=Path, default=Path("rag_index"))
    query_parser.add_argument("--top-k", type=int, default=3)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    try:
        if args.command == "index":
            files = collect_files(args.inputs)
            if not files:
                raise FileNotFoundError("No supported PDF, TXT, or DOCX files were found.")
            build_index(files, args.index_dir, args.chunk_size, args.overlap)
        else:
            query_index(args.question, args.index_dir, args.top_k)
    except Exception as error:
        print(f"Error: {error}", file=sys.stderr)
        raise SystemExit(1)


if __name__ == "__main__":
    main()
